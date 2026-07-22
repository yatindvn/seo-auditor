"""
report.py — Aggregation of audit results into a health score + exports.
"""

from __future__ import annotations

import csv
import json
import os
from collections import Counter
from datetime import datetime
from typing import Dict, List


def compute_health_score(all_page_issues: Dict[str, List[Dict]], total_pages: int) -> Dict:
    """0-100 score. Deduct weighted points per issue severity, normalized by page count."""
    if total_pages == 0:
        return {"score": 0, "grade": "N/A"}

    sev_weight = {"critical": 5, "warning": 2, "info": 0.5}
    total_deduction = 0
    counts = Counter()
    for issues in all_page_issues.values():
        for iss in issues:
            total_deduction += sev_weight.get(iss["severity"], 1)
            counts[iss["severity"]] += 1

    avg_deduction_per_page = total_deduction / total_pages
    score = max(0, 100 - avg_deduction_per_page * 4)
    score = round(min(score, 100), 1)

    if score >= 90:
        grade = "A"
    elif score >= 75:
        grade = "B"
    elif score >= 60:
        grade = "C"
    elif score >= 40:
        grade = "D"
    else:
        grade = "F"

    return {
        "score": score,
        "grade": grade,
        "critical_issues": counts.get("critical", 0),
        "warning_issues": counts.get("warning", 0),
        "info_issues": counts.get("info", 0),
    }


def build_executive_summary(crawler, all_page_issues, site_wide, duplicates, redirects, broken_links, recommendations) -> Dict:
    health = compute_health_score(all_page_issues, len(crawler.results))
    status_counts = Counter(r.status_code for r in crawler.results.values())

    return {
        "audit_date": datetime.utcnow().isoformat() + "Z",
        "start_url": crawler.start_url,
        "pages_crawled": len(crawler.results),
        "health_score": health,
        "status_code_breakdown": dict(status_counts),
        "orphan_pages": site_wide["orphan_count"],
        "broken_links": len(broken_links),
        "redirect_chains": len(redirects["redirect_chains"]),
        "redirect_loops": len(redirects["redirect_loops"]),
        "duplicate_titles": len(duplicates["duplicate_titles"]),
        "duplicate_meta_descriptions": len(duplicates["duplicate_meta_descriptions"]),
        "duplicate_content_groups": len(duplicates["duplicate_content"]),
        "top_recommendations": recommendations[:10],
        "robots_txt_found": crawler.robots_txt_content is not None,
        "sitemaps_found": crawler.sitemaps_found,
        "sitemap_url_count": len(crawler.sitemap_urls),
    }


def build_page_level_report(crawler, all_page_issues, page_meta: Dict[str, Dict]) -> List[Dict]:
    rows = []
    for url, result in crawler.results.items():
        meta = page_meta.get(url, {})
        issues = all_page_issues.get(url, [])
        rows.append({
            "url": url,
            "status_code": result.status_code,
            "depth": result.depth,
            "response_time_ms": result.response_time_ms,
            "title": meta.get("title"),
            "meta_description": meta.get("meta_description"),
            "h1": meta.get("h1"),
            "word_count": meta.get("word_count"),
            "canonical": meta.get("canonical"),
            "critical_issues": sum(1 for i in issues if i["severity"] == "critical"),
            "warning_issues": sum(1 for i in issues if i["severity"] == "warning"),
            "info_issues": sum(1 for i in issues if i["severity"] == "info"),
            "issue_codes": ", ".join(sorted({i["code"] for i in issues})),
        })
    return rows


# --------------------------------------------------------------------- export
def export_json(output_path: str, data: Dict):
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, default=str)


def export_csv(output_path: str, rows: List[Dict]):
    if not rows:
        with open(output_path, "w") as f:
            f.write("")
        return
    fieldnames = list(rows[0].keys())
    with open(output_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def export_excel(output_path: str, sheets: Dict[str, List[Dict]]):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)

    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for sheet_name, rows in sheets.items():
        ws = wb.create_sheet(title=sheet_name[:31])
        if not rows:
            ws.append(["No data"])
            continue
        headers = list(rows[0].keys())
        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
        for row in rows:
            ws.append([str(row.get(h, "")) if not isinstance(row.get(h), (int, float, type(None))) else row.get(h) for h in headers])
        for col_idx, header in enumerate(headers, 1):
            max_len = max([len(str(header))] + [len(str(r.get(header, ""))) for r in rows[:200]])
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 60)
        ws.freeze_panes = "A2"

    wb.save(output_path)


def export_pdf_summary(output_path: str, executive_summary: Dict):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

    doc = SimpleDocTemplate(output_path, pagesize=letter, topMargin=0.6 * inch, bottomMargin=0.6 * inch)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("TitleBig", parent=styles["Title"], fontSize=20)
    story = []

    story.append(Paragraph("SEO Audit Report", title_style))
    story.append(Paragraph(f"URL: {executive_summary['start_url']}", styles["Normal"]))
    story.append(Paragraph(f"Date: {executive_summary['audit_date']}", styles["Normal"]))
    story.append(Spacer(1, 16))

    health = executive_summary["health_score"]
    story.append(Paragraph(f"Health Score: {health['score']} / 100 (Grade {health['grade']})", styles["Heading2"]))
    story.append(Spacer(1, 8))

    summary_rows = [
        ["Pages Crawled", executive_summary["pages_crawled"]],
        ["Critical Issues", health.get("critical_issues", 0)],
        ["Warning Issues", health.get("warning_issues", 0)],
        ["Info Issues", health.get("info_issues", 0)],
        ["Orphan Pages", executive_summary["orphan_pages"]],
        ["Broken Links", executive_summary["broken_links"]],
        ["Redirect Chains", executive_summary["redirect_chains"]],
        ["Redirect Loops", executive_summary["redirect_loops"]],
        ["Duplicate Titles", executive_summary["duplicate_titles"]],
        ["Duplicate Meta Descriptions", executive_summary["duplicate_meta_descriptions"]],
    ]
    t = Table(summary_rows, colWidths=[260, 200])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F2937")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.whitesmoke, colors.white]),
    ]))
    story.append(t)
    story.append(Spacer(1, 20))

    story.append(Paragraph("Top Recommendations", styles["Heading2"]))
    for i, rec in enumerate(executive_summary["top_recommendations"], 1):
        story.append(Paragraph(
            f"{i}. [{rec['severity'].upper()}] {rec['recommendation']} "
            f"({rec['affected_pages']} page(s) affected)",
            styles["Normal"],
        ))
        story.append(Spacer(1, 4))

    doc.build(story)


def export_html_report(output_path: str, executive_summary: Dict, page_rows: List[Dict]):
    health = executive_summary["health_score"]
    grade_colors = {"A": "#16a34a", "B": "#65a30d", "C": "#ca8a04", "D": "#ea580c", "F": "#dc2626"}
    color = grade_colors.get(health["grade"], "#6b7280")

    rec_rows_html = "".join(
        f"<tr><td>{r['severity'].upper()}</td><td>{r['recommendation']}</td>"
        f"<td>{r['affected_pages']}</td></tr>"
        for r in executive_summary["top_recommendations"]
    )
    page_rows_html = "".join(
        f"<tr><td>{r['url']}</td><td>{r['status_code']}</td><td>{r['title'] or ''}</td>"
        f"<td>{r['critical_issues']}</td><td>{r['warning_issues']}</td><td>{r['info_issues']}</td></tr>"
        for r in page_rows[:500]
    )

    html = f"""<!DOCTYPE html>
<html><head><meta charset="utf-8"><title>SEO Audit Report</title>
<style>
body {{ font-family: -apple-system, Arial, sans-serif; margin: 40px; color: #111827; background: #f9fafb; }}
h1 {{ margin-bottom: 4px; }}
.meta {{ color: #6b7280; margin-bottom: 24px; }}
.score {{ display: inline-block; background: {color}; color: white; font-size: 32px; font-weight: bold;
          padding: 16px 28px; border-radius: 12px; }}
table {{ border-collapse: collapse; width: 100%; background: white; margin-top: 16px; }}
th, td {{ text-align: left; padding: 8px 12px; border-bottom: 1px solid #e5e7eb; font-size: 13px; }}
th {{ background: #1f2937; color: white; }}
tr:hover {{ background: #f3f4f6; }}
.cards {{ display: flex; gap: 16px; margin: 20px 0; flex-wrap: wrap; }}
.card {{ background: white; border-radius: 10px; padding: 16px 20px; box-shadow: 0 1px 3px rgba(0,0,0,.1); min-width: 140px; }}
.card .num {{ font-size: 24px; font-weight: bold; }}
.card .label {{ color: #6b7280; font-size: 12px; text-transform: uppercase; }}
</style></head><body>
<h1>SEO Audit Report</h1>
<div class="meta">{executive_summary['start_url']} &middot; {executive_summary['audit_date']}</div>
<div class="score">{health['score']} / 100 &middot; Grade {health['grade']}</div>

<div class="cards">
  <div class="card"><div class="num">{executive_summary['pages_crawled']}</div><div class="label">Pages Crawled</div></div>
  <div class="card"><div class="num">{health.get('critical_issues', 0)}</div><div class="label">Critical Issues</div></div>
  <div class="card"><div class="num">{health.get('warning_issues', 0)}</div><div class="label">Warnings</div></div>
  <div class="card"><div class="num">{executive_summary['orphan_pages']}</div><div class="label">Orphan Pages</div></div>
  <div class="card"><div class="num">{executive_summary['broken_links']}</div><div class="label">Broken Links</div></div>
  <div class="card"><div class="num">{executive_summary['duplicate_titles']}</div><div class="label">Duplicate Titles</div></div>
</div>

<h2>Top Recommendations</h2>
<table><tr><th>Severity</th><th>Recommendation</th><th>Pages Affected</th></tr>{rec_rows_html}</table>

<h2>Pages ({len(page_rows)})</h2>
<table><tr><th>URL</th><th>Status</th><th>Title</th><th>Critical</th><th>Warning</th><th>Info</th></tr>{page_rows_html}</table>

</body></html>"""

    with open(output_path, "w", encoding="utf-8") as f:
        f.write(html)
